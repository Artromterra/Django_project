"""The module responsible for importing data from xlsx files."""

from typing import Generator, Dict, Any, Tuple, Literal, Set
from datetime import datetime
from copy import deepcopy

from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.worksheet import Worksheet

from services.importing.importers.base import BaseImporter
from services.importing.excs import NotFoundData, NotFoundField


class XLSXImporter(BaseImporter):
    """
    The class responsible for importing entities from xlsx files.

    The import file must have the xlsx extension,
    and the categories, products, sellers, and product_seller pages
    must have tables with import data.
    The first line should contain the names of the fields
    that correspond to the fields of the models.
    Fields that link to fields in other tables must have foreign keys.
    Cells with no value are interpreted as null;
    cells with an explicitly specified empty string ("")
    are interpreted as an empty string.
    """

    def __init__(self, filename: str, logger_name: str = "main.services.importing.importers"):
        """
        Init class.

        :param filename: Import file (abs path)
        :param logger_name: The name of the logger.
        It is necessary to create separate lags when importing multiple files in parallel.
        """
        super().__init__(filename, logger_name)
        self.__wb: Workbook = load_workbook(filename, read_only=True)

        self.import_params["Product"].update(sheet="products")
        self.import_params["Seller"].update(sheet="sellers")
        self.import_params["Category"].update(sheet="categories")
        self.import_params["ProductSeller"].update(sheet="product_seller")

    def __check_sheet_exists(
        self,
        model: Literal["Category", "Product", "Seller", "ProductSeller"],
    ) -> str:
        """
        Check that sheet with data for model exists.

        :raise NotFoundData: If sheet not exists.
        :return: Sheet name.
        """
        sheet_name = self.import_params[model]["sheet"]

        if sheet_name not in self.__wb:
            raise NotFoundData(
                model=model,
                msg=f"Not found sheet {sheet_name} in {self.filename}"
            )

        return sheet_name

    def __check_headers(
        self,
        model: Literal["Category", "Product", "Seller", "ProductSeller"],
        headers: Tuple[str, ...]
    ) -> Set[str]:
        """
        Check that the headers match the fields of the model.

        :raise NotFoundField: If not found all fields in sheet.
        :return: set of model fields.
        """
        headers_set: Set[str] = set(headers)

        model_fields = set(deepcopy(self.import_params[model]["fields"]))
        if not model_fields.issubset(headers_set):
            raise NotFoundField(
                model=model,
                fields=model_fields - (model_fields & headers_set),
                msg=f"Not found headers in {self.filename}"
            )

        return model_fields

    def __match_headers_and_col_nums(
        self,
        model: Literal["Category", "Product", "Seller", "ProductSeller"],
        model_fields: Set[str],
        headers: Tuple[str, ...],
    ) -> None:
        """
        Match headings and column numbers.

        :param model: Model name.
        :param model_fields: Set of model fields.
        :param headers: Tuple of headers.
        """
        self.import_params[model]["cols"] = dict()
        for num, header in enumerate(headers):
            if header in model_fields:
                self.import_params[model]["cols"][header] = num
                model_fields.remove(header)

    def __pre_import(
        self,
        model: Literal["Category", "Product", "Seller", "ProductSeller"],
    ) -> Generator[tuple[str | float | datetime | None, ...], None, None]:
        """
        Check data before import.

        :param model: Model name.

        :raise NotFoundData: if not found sheet with sheet_name.
        :raise NotFoundField: if not found some fields in headers.
        :return: Generator of rows (without headers).
        """
        # check that sheet exists
        self.logger.info("Check that sheet with %s data exists.", model)
        sheet_name: str = self.__check_sheet_exists(model)
        ws: Worksheet = self.__wb[sheet_name]

        rows_iter = ws.iter_rows(values_only=True)
        headers: Tuple[str, ...] = next(rows_iter)

        self.logger.info("Check that headers are correct.")
        # check that the headers match the fields of the model
        model_fields: Set[str] = self.__check_headers(model, headers)

        self.logger.debug("Match headers and col numbers.")
        # compare the column number to each heading
        self.__match_headers_and_col_nums(model, model_fields, headers)
        return rows_iter

    def __import_data(
        self,
        model: Literal["Category", "Product", "Seller", "ProductSeller"],
    ) -> Generator[Dict[str, Any], None, None]:
        """Import model data."""
        # check data before importing
        rows_iter = self.__pre_import(model)

        # yield dict with data
        for row in rows_iter:
            model_data: Dict[str, Any] = {
                header: row[col]
                for header, col in self.import_params[model]["cols"].items()
            }
            self.logger.debug("Yield dict with %s data", model)
            yield model_data


    def import_categories(self) -> Generator[Dict[str, Any], None, None]:
        """Import categories."""
        self.logger.info("Start importing categories.")
        try:
            for category_data in self.__import_data("Category"):
                yield category_data
        except (NotFoundData, NotFoundField) as exc:
            self.logger.warning("Can't import categories from %s\n%s", self.filename, str(exc))
        else:
            self.logger.info("The categories have been successfully imported.")


    def import_products(self) -> Generator[Dict[str, Any], None, None]:
        """Import products."""
        self.logger.info("Start importing products.")
        try:
            for product_data in self.__import_data("Product"):
                yield product_data
        except (NotFoundData, NotFoundField) as exc:
            self.logger.warning("Can't import products from %s\n%s", self.filename, str(exc))
        else:
            self.logger.info("The products have been successfully imported.")

    def import_sellers(self) -> Generator[Dict[str, Any], None, None]:
        """Import sellers."""
        self.logger.info("Start importing sellers.")
        try:
            for seller_data in self.__import_data("Seller"):
                yield seller_data
        except (NotFoundData, NotFoundField) as exc:
            self.logger.warning("Can't import sellers from %s\n%s", self.filename, str(exc))
        else:
            self.logger.info("The sellers have been successfully imported.")

    def import_products_sellers(self) -> Generator[Dict[str, Any], None, None]:
        """Import ProductSeller instances."""
        self.logger.info("Start importing ProductSeller instances.")
        try:
            for product_seller_data in self.__import_data("ProductSeller"):
                yield product_seller_data
        except (NotFoundData, NotFoundField) as exc:
            self.logger.warning("Can't import ProductSeller from %s\n%s", self.filename, str(exc))
        else:
            self.logger.info("The ProductSeller have been successfully imported.")
